"""Offline verification of the job_scraper fixes.

No network: a FakeSession serves a synthetic multi-page career site, so the
assertions are about the scraper's own logic (caps, pagination, counting,
column mapping) rather than about any live website.
"""
import csv, io, json, os, shutil, socket, sys, tempfile, time

ROOT = os.path.dirname(os.path.abspath(__file__))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from job_scraper import company, config, netcache, pipeline, websearch
from job_scraper.urlutils import normalize_website

PASS, FAIL = [], []


def check(label, condition, detail=""):
    (PASS if condition else FAIL).append(label)
    print("%s %s%s" % ("PASS" if condition else "FAIL", label,
                       (" -- " + str(detail)) if detail else ""))


# ---------------------------------------------------------------- fixture site
PAGES = 5
JOBS_PER_PAGE = 30
TOTAL_JOBS = PAGES * JOBS_PER_PAGE
HOST = "https://fixture-co.example"


def _jsonld(page):
    postings = []
    for i in range(JOBS_PER_PAGE):
        n = (page - 1) * JOBS_PER_PAGE + i + 1
        postings.append({
            "@context": "https://schema.org",
            "@type": "JobPosting",
            "title": "Engineer %03d" % n,
            "description": "<p>Responsibilities. Apply now. 5 years of experience.</p>",
            "datePosted": "2026-08-01",
            "employmentType": "FULL_TIME",
            "hiringOrganization": {"@type": "Organization", "name": "Fixture Co"},
            "jobLocation": {"@type": "Place",
                            "address": {"@type": "PostalAddress", "addressLocality": "Berlin"}},
            "url": "%s/careers/job/%d" % (HOST, n),
        })
    return json.dumps(postings)


def career_page(page):
    next_link = ('<a href="/careers?page=%d">Next</a>' % (page + 1)) if page < PAGES else ""
    return (
        "<html><head><title>Careers at Fixture Co</title>"
        '<script type="application/ld+json">%s</script></head><body><main>'
        "<h1>Open positions</h1>"
        "<p>Apply now. Job description, responsibilities, qualifications.</p>"
        "%s</main></body></html>" % (_jsonld(page), next_link)
    )


HOMEPAGE = (
    "<html><body><h1>Fixture Co</h1>"
    + "".join('<a href="/careers?page=1">Careers</a>')
    + "".join('<a href="/careers/team-%d">Careers team %d</a>' % (i, i) for i in range(1, 21))
    + "</body></html>"
)


class FakeResponse:
    def __init__(self, url, text, status=200, ctype="text/html"):
        self.url, self.text, self.status_code = url, text, status
        self.headers = {"Content-Type": ctype}

    def json(self):
        return json.loads(self.text)


class FakeSession:
    """Serves the fixture site and records every fetch."""

    def __init__(self):
        self.fetched = []

    def _body(self, url):
        u = url.split("#")[0]
        if u.rstrip("/") in (HOST, HOST + "/"):
            return HOMEPAGE
        if "/careers?page=" in u:
            try:
                page = int(u.split("page=")[1].split("&")[0])
            except ValueError:
                page = 1
            return career_page(page) if 1 <= page <= PAGES else None
        if u.endswith("/careers") or u.endswith("/careers/"):
            return career_page(1)
        if "/careers/job/" in u:
            n = u.rsplit("/", 1)[-1]
            return ("<html><head><script type=\"application/ld+json\">%s</script>"
                    "</head><body>Apply</body></html>"
                    % json.dumps({"@type": "JobPosting", "title": "Engineer %s" % n,
                                  "url": u, "description": "Detail"}))
        return None

    def fetch(self, url, timeout=None, method="GET", retries=None, **kw):
        self.fetched.append(url)
        body = self._body(url)
        if body is None:
            return None
        return FakeResponse(url, body)

    def fetch_text(self, url, timeout=None, **kw):
        r = self.fetch(url, timeout=timeout, **kw)
        return r.text if r else None

    def fetch_json(self, url, timeout=None, **kw):
        return None

    def post_json(self, url, payload, timeout=None):
        return None

    def head_ok(self, url, timeout=12):
        return self._body(url) is not None


# ------------------------------------------------------- 1. input column mapping
XLSX = os.path.join(ROOT, "smoke_test_input.xlsx")
if os.path.exists(XLSX):
    rows = pipeline.read_companies(XLSX)
    check("read_companies returns every data row", len(rows) > 900, len(rows))
    first = rows[0]
    check("company_name column mapped (not KEYID hash)",
          first[0] == "ROADS AGENCY LIMPOPO (PTY) LTD", first[0])
    check("country column mapped", first[2] == "South Africa", first[2])
    carolina = [r for r in rows if r[0] == "CAROLINA IMAGING INC"]
    check("website column mapped from WEBSITE header",
          bool(carolina) and carolina[0][1] == "www.carolinamri.com",
          carolina[0] if carolina else None)
    check("no company name ever lands in the website field",
          not any(" " in (r[1] or "") for r in rows))
    with_site = sum(1 for r in rows if r[1])
    print("     rows with a real website value: %d/%d" % (with_site, len(rows)))

# headerless sheet keeps the old positional behaviour
import openpyxl
tmpdir = tempfile.mkdtemp()
hl = os.path.join(tmpdir, "headerless.xlsx")
wb = openpyxl.Workbook(); ws = wb.active
ws.append(["Acme Ltd", "acme.example", "USA"]); ws.append(["Beta Inc", "beta.example", "UK"])
wb.save(hl)
rows = pipeline.read_companies(hl)
check("headerless sheet falls back to positional layout",
      rows == [("Acme Ltd", "acme.example", "USA"), ("Beta Inc", "beta.example", "UK")], rows)

# ------------------------------------------------------- 2. hostname validation
check("company name rejected as website", normalize_website("ROADS AGENCY LIMPOPO (PTY) LTD") == "")
check("real domain accepted",
      normalize_website("www.carolinamri.com") == "https://www.carolinamri.com")
check("IDN domain accepted", normalize_website("müller.de") == "https://müller.de")

# ------------------------------------------------------- 3. DNS negative caching
netcache.install()
bogus = "roads agency limpopo (pty) ltd"
t0 = time.time()
for _ in range(50):
    try:
        socket.getaddrinfo(bogus, 443)
    except socket.gaierror:
        pass
elapsed = time.time() - t0
check("50 lookups of a non-hostname take < 0.2s (negative cache)", elapsed < 0.2,
      "%.4fs" % elapsed)
check("negative cache recorded hits", netcache.stats()["negative_hits"] >= 49,
      netcache.stats())

# ------------------------------------------------------- 4. extraction, no caps
details = company.process_company_details(("Fixture Co", HOST, "Germany"),
                                          session=FakeSession(), enable_search=False)
n = details.get("num_positions")
check("process_company_details returns num_positions", n is not None, n)
check("all %d postings across %d paginated pages extracted" % (TOTAL_JOBS, PAGES),
      n == TOTAL_JOBS, "got %s" % n)
check("num_positions equals len(jobs)", n == len(details["jobs"]))
check("status ok", details["status"] == "ok", details["status"])
check("career page discovered", bool(details["career_page_url"]), details["career_page_url"])
titles = {j.get("job_title") for j in details["jobs"]}
check("last page's postings present (no page-1 truncation)",
      "Engineer %03d" % TOTAL_JOBS in titles)
check("no 200-job ceiling", n > 200 or TOTAL_JOBS <= 200, n)

# ------------------------------------------------------- 5. search engine chain
calls = {"a": 0, "b": 0}


def engine_down(session, query):
    calls["a"] += 1
    return None            # engine failure


def engine_up(session, query):
    calls["b"] += 1
    return ["https://found.example/careers"]


saved = websearch._ENGINES
websearch._ENGINES = (("down", engine_down), ("up", engine_up))
websearch._QUERY_CACHE.clear(); websearch._ENGINE_STATE.clear()
out = websearch.web_search(FakeSession(), "fixture co careers")
check("failing engine falls through to the next", out == ["https://found.example/careers"], out)
check("failed engine retried SEARCH_MAX_ATTEMPTS times",
      calls["a"] == config.SEARCH_MAX_ATTEMPTS, calls["a"])
before = calls["b"]
websearch.web_search(FakeSession(), "fixture co careers")
check("identical query served from cache", calls["b"] == before, calls["b"])
for i in range(config.SEARCH_ENGINE_MAX_FAILURES + 1):
    websearch._QUERY_CACHE.clear()
    websearch.web_search(FakeSession(), "cb probe %d" % i)
state = websearch.engine_status().get("down", {})
check("failing engine benched by circuit breaker", bool(state.get("benched_until")), state)
websearch._ENGINES = saved
websearch._QUERY_CACHE.clear(); websearch._ENGINE_STATE.clear()

# ------------------------------------------------------- 6. CSV schema migration
legacy = os.path.join(tmpdir, "ajay_legacy.csv")
legacy_cols = [c for c in config.OUTPUT_COLUMNS if c != "num_positions"]
with io.open(legacy, "w", encoding="utf-8-sig", newline="") as f:
    writer = csv.writer(f)
    writer.writerow(legacy_cols)
    for i in range(5):
        writer.writerow(["Co%d" % i] + [""] * (len(legacy_cols) - 1))
with io.open(legacy, encoding="utf-8-sig") as f:
    before_cols = next(csv.reader(f))
    before_rows = sum(1 for _ in f)
check("synthetic legacy CSV lacks num_positions", "num_positions" not in before_cols)
pipeline.CsvWriter(legacy, config.OUTPUT_COLUMNS)
with io.open(legacy, encoding="utf-8-sig") as f:
    after = list(csv.DictReader(f))
cols = list(after[0].keys()) if after else []
check("migrated file has num_positions column", "num_positions" in cols)
check("migration preserved every row", len(after) == before_rows,
      "%d -> %d" % (before_rows, len(after)))
check("migration left existing values intact",
      bool(after) and after[0]["company_name"] == "Co0", after[0] if after else None)

# a file already on the current schema must be left alone
current = os.path.join(tmpdir, "current.csv")
pipeline.CsvWriter(current, config.OUTPUT_COLUMNS)
before_mtime = os.path.getmtime(current)
time.sleep(0.01)
pipeline.CsvWriter(current, config.OUTPUT_COLUMNS)
check("matching schema is not rewritten", os.path.getmtime(current) == before_mtime)

# ------------------------------------------------------- 7. end-to-end pipeline
inp = os.path.join(tmpdir, "in.xlsx")
wb = openpyxl.Workbook(); ws = wb.active
ws.append(["KEYID", "COMPANY_NAME", "COUNTRY", "ENTITY_TYPE", "WEBSITE"])
ws.append(["hash1", "Fixture Co", "Germany", "Independent", "fixture-co.example"])
ws.append(["hash2", "Nowhere GmbH", "Germany", "Independent", None])
wb.save(inp)
out_csv = os.path.join(tmpdir, "out.csv")
pipeline.ScrapeSession = FakeSession           # inject the offline session
counters = pipeline.run(input_file=inp, output_file=out_csv, workers=2,
                        resume=False, enable_search=False, quiet=True)
with io.open(out_csv, encoding="utf-8-sig") as f:
    written = list(csv.DictReader(f))
check("pipeline wrote rows", len(written) > 0, len(written))
check("num_positions in written header", "num_positions" in written[0])
fixture_rows = [r for r in written if r["company_name"] == "Fixture Co"]
check("every Fixture Co row carries num_positions=%d" % TOTAL_JOBS,
      bool(fixture_rows) and all(r["num_positions"] == str(TOTAL_JOBS) for r in fixture_rows),
      {r["num_positions"] for r in fixture_rows})
check("one CSV row per position", len(fixture_rows) == TOTAL_JOBS, len(fixture_rows))
zero_rows = [r for r in written if r["company_name"] == "Nowhere GmbH"]
check("zero-job company gets num_positions=0 (not blank)",
      bool(zero_rows) and zero_rows[0]["num_positions"] == "0",
      zero_rows[0]["num_positions"] if zero_rows else None)
check("counters total jobs matches", counters.get("jobs") == TOTAL_JOBS, counters)



# ------------------------- 8. external-ATS handoff + career-hub traversal
# Regression test for vumc.org/careers: the hub page has no postings, the real
# jobs are on vumc.wd1.myworkdayjobs.com, and three same-domain branches carry
# more.  Previously reported "Career Page Found - Extraction Unsupported".
HOME = "https://www.vumc.org"
HUB = HOME + "/careers"
WD = "https://vumc.wd1.myworkdayjobs.com/vumccareers"
BRANCHES = {"/careers/nursing": 12, "/careers/physician": 7, "/careers/allied-health": 5}
WD_TOTAL = 40


class R:
    def __init__(s, u, t, ct="text/html"):
        s.url, s.text, s.status_code, s.headers = u, t, 200, {"Content-Type": ct}
    def json(s):
        return json.loads(s.text)


def branch_html(path, n):
    ld = json.dumps([{"@type": "JobPosting", "title": "%s role %d" % (path.strip('/').split('/')[-1], i),
                      "url": HOME + path + "/j%d" % i, "description": "Apply. Responsibilities."}
                     for i in range(n)])
    return ('<html><head><title>Careers</title>'
            '<script type="application/ld+json">%s</script></head>'
            '<body><main><h1>Open positions</h1><p>Apply now</p></main></body></html>' % ld)


HUB_HTML = """<html><head><title>Work for Vanderbilt Health | Careers</title></head><body><main>
<h1>Work for Vanderbilt Health</h1>
<a href="/careers">Careers Home</a>
<a href="https://vumc.wd1.myworkdayjobs.com/vumccareers">Search Jobs</a>
<a href="/careers/nursing">Nursing Careers</a>
<a href="/careers/allied-health">Allied Health Careers</a>
<a href="/careers/physician">Physician Careers</a>
<a href="/careers/how-to-apply">How to Apply</a>
<a href="/careers/contact">Contact Us</a>
</main></body></html>"""


class FS:
    def __init__(s): s.calls = []
    def _body(s, url):
        u = url.split("#")[0].rstrip("/")
        if u in (HOME, HOME + "/index.html"):
            return ('<html><body><a href="/careers">Careers</a></body></html>', "text/html")
        if u == HUB:
            return (HUB_HTML, "text/html")
        for path, n in BRANCHES.items():
            if u == HOME + path:
                return (branch_html(path, n), "text/html")
        if u in (HOME + "/careers/how-to-apply", HOME + "/careers/contact"):
            return ("<html><body>How to apply. No postings here.</body></html>", "text/html")
        if "/wday/cxs/vumc/vumccareers/jobs" in u:
            return None  # handled in post_json
        if u.startswith(WD):
            return ("<html><body>Workday board shell</body></html>", "text/html")
        return None
    def fetch(s, url, timeout=None, method="GET", retries=None, **k):
        s.calls.append(url)
        body = s._body(url)
        return R(url, body[0], body[1]) if body else None
    def fetch_text(s, url, timeout=None, **k):
        r = s.fetch(url, timeout=timeout)
        return r.text if r else None
    def fetch_json(s, url, timeout=None, **k):
        return None
    def post_json(s, url, payload, timeout=None):
        s.calls.append("POST " + url)
        if "/wday/cxs/vumc/vumccareers/jobs" not in url:
            return None          # wrong tenant -> the old bug's 404
        offset = payload.get("offset", 0)
        limit = payload.get("limit", 20)
        items = [{"title": "Workday Job %03d" % i,
                  "externalPath": "/job/%d" % i,
                  "postedOn": "Posted Today",
                  "locationsText": "Nashville, TN"}
                 for i in range(offset, min(offset + limit, WD_TOTAL))]
        return {"total": WD_TOTAL, "jobPostings": items}
    def head_ok(s, *a, **k): return True


vumc_session = FS()
details = company.process_company_details(("Vanderbilt University Medical Center", HOME, "USA"),
                                          session=vumc_session, enable_search=False)

print("num_positions       :", details["num_positions"])
print("career_page_url     :", details["career_page_url"])
print("discovery method    :", details["career_page_discovery_method"])
print("source              :", details["source"])

titles = [j.get("job_title", "") for j in details["jobs"]]
wd = [t for t in titles if t.startswith("Workday Job")]
nursing = [t for t in titles if t.startswith("nursing role")]
phys = [t for t in titles if t.startswith("physician role")]
allied = [t for t in titles if t.startswith("allied-health role")]
print("\nworkday board jobs  : %d (expect %d)" % (len(wd), WD_TOTAL))
print("nursing branch      : %d (expect %d)" % (len(nursing), BRANCHES['/careers/nursing']))
print("physician branch    : %d (expect %d)" % (len(phys), BRANCHES['/careers/physician']))
print("allied branch       : %d (expect %d)" % (len(allied), BRANCHES['/careers/allied-health']))

check("external ATS handoff: Workday board reached",
      len(wd) == WD_TOTAL, "%d of %d" % (len(wd), WD_TOTAL))
check("career-hub branches walked",
      len(nursing) == BRANCHES["/careers/nursing"] and
      len(phys) == BRANCHES["/careers/physician"] and
      len(allied) == BRANCHES["/careers/allied-health"],
      "%d/%d/%d" % (len(nursing), len(phys), len(allied)))
check("hub company no longer 'Extraction Unsupported'",
      details["status"] == "ok", details["status"])
check("num_positions counts hub + external board",
      details["num_positions"] == WD_TOTAL + sum(BRANCHES.values()),
      details["num_positions"])
check("correct Workday tenant endpoint called",
      any("/wday/cxs/vumc/vumccareers/jobs" in c for c in vumc_session.calls))
check("career_page_url points at the board that has the jobs",
      details["career_page_url"] == WD, details["career_page_url"])

# ------------------- 9. vendor coverage + iframe-embedded board extraction
from job_scraper.ats import detect_ats_in_url as _detect

VENDOR_URLS = [
    ("https://jobs.ashbyhq.com/notion", "ashby"),
    ("https://ats.rippling.com/acme-corp/jobs", "rippling"),
    ("https://acme.csod.com/ux/ats/careersite/4/home", "cornerstone"),
    ("https://sub.dayforcehcm.com/CandidatePortal/en-US/acme", "dayforce"),
    ("https://acme.eightfold.ai/careers", "eightfold"),
    ("https://recruiting.paylocity.com/recruiting/jobs/All/abc-123/Acme", "paylocity"),
    ("https://acme.paycomonline.net/v4/ats/web.php/jobs", "paycom"),
    ("https://recruiting.ultipro.com/ACM1234ACME/JobBoard/x/", "ukg"),
    ("https://workforcenow.adp.com/mascsr/default/mdf/recruitment/x.html", "adp"),
    ("https://acme.myisolved.com/careers", "isolved"),
    ("https://acme.homerun.co/", "homerun"),
    ("https://acme.keka.com/careers/", "keka"),
    ("https://acme.darwinbox.in/ms/candidate/careers", "darwinbox"),
    ("https://acme.jobylon.com/jobs/", "jobylon"),
    ("https://acme.factorialhr.com/job_posting", "factorial"),
    ("https://acme.wd5.myworkdaysite.com/en-US/AcmeCareers", "workday"),
    ("https://careers.smartrecruiters.com/AcmeInc", "smartrecruiters"),
    ("https://job-boards.greenhouse.io/acme", "greenhouse"),
]
wrong = [(u, _detect(u)[0], want) for u, want in VENDOR_URLS if _detect(u)[0] != want]
check("all %d vendor URL shapes detected" % len(VENDOR_URLS), not wrong, wrong[:3])
check("a plain careers page is not misdetected as an ATS",
      _detect("https://www.example.com/careers")[0] is None)

ASHBY_HOME = "https://www.acmebio.com"
ASHBY_HUB = ASHBY_HOME + "/careers"
ASHBY_N = 25


class AshbyIframeSession:
    """Board embedded via <iframe> only -- no <a href> to find."""

    def __init__(self):
        self.calls = []

    def _body(self, url):
        u = url.split("#")[0].rstrip("/")
        if u == ASHBY_HOME:
            return '<html><body><a href="/careers">Careers</a></body></html>'
        if u == ASHBY_HUB:
            return ('<html><head><title>Careers at AcmeBio</title></head><body><main>'
                    '<h1>Open roles</h1>'
                    '<iframe src="https://jobs.ashbyhq.com/acmebio?embed=js"></iframe>'
                    '</main></body></html>')
        if u.startswith("https://jobs.ashbyhq.com/acmebio"):
            return '<html><body>Ashby board shell</body></html>'
        return None

    def fetch(self, url, timeout=None, method="GET", retries=None, **kw):
        self.calls.append(url)
        body = self._body(url)
        return FakeResponse(url, body) if body else None

    def fetch_text(self, url, timeout=None, **kw):
        r = self.fetch(url, timeout=timeout)
        return r.text if r else None

    def fetch_json(self, url, timeout=None, **kw):
        self.calls.append("JSON " + url)
        if "api.ashbyhq.com/posting-api/job-board/acmebio" in url:
            return {"jobs": [{"id": "id%d" % i, "title": "Scientist %02d" % i,
                              "isListed": True, "location": "Boston, MA",
                              "jobUrl": "https://jobs.ashbyhq.com/acmebio/id%d" % i,
                              "employmentType": "FullTime", "department": "R&D",
                              "publishedAt": "2026-08-01T00:00:00Z",
                              "descriptionHtml": "<p>Apply now.</p>",
                              "compensation": {"compensationTierSummary": "$120K - $150K"}}
                             for i in range(ASHBY_N)]}
        return None

    def post_json(self, url, payload, timeout=None):
        return None

    def head_ok(self, url, timeout=12):
        return True


ashby_session = AshbyIframeSession()
ashby = company.process_company_details(("AcmeBio Inc", ASHBY_HOME, "USA"),
                                        session=ashby_session, enable_search=False)
check("iframe-embedded board discovered (no anchor present)",
      ashby["num_positions"] == ASHBY_N, ashby["num_positions"])
check("Ashby posting API called with the right board",
      any("posting-api/job-board/acmebio" in c for c in ashby_session.calls))
check("Ashby compensation captured",
      bool(ashby["jobs"] and ashby["jobs"][0].get("salary")),
      ashby["jobs"][0].get("salary") if ashby["jobs"] else None)

# ------------------------------- 10. derived field enrichment
# Every ATS API parser returns title/location/url and little else; the rest of
# OUTPUT_COLUMNS is derived here.  Previously that derivation ran only inside
# parse_generic, so every ATS-sourced row exported blank.
from job_scraper.fields import enrich_job, parse_date, salary_range_from_text
import datetime as _dt

_today = _dt.datetime.now(_dt.timezone.utc).strftime("%Y-%m-%d")
check("relative date 'Posted Today' resolved", parse_date("Posted Today") == _today,
      parse_date("Posted Today"))
check("relative date 'Posted 30+ Days Ago' resolved",
      bool(parse_date("Posted 30+ Days Ago")), parse_date("Posted 30+ Days Ago"))
check("ISO dates still parsed", parse_date("2026-08-01") == "2026-08-01")

check("salary range without a keyword", salary_range_from_text("$70,000 - $90,000") != "")
check("hourly rate detected", salary_range_from_text("$32.50 per hour") != "")
check("unrelated price not treated as salary",
      salary_range_from_text("Our product costs $49 per month") == "")
check("benefits text not treated as salary",
      salary_range_from_text("We have 13 paid state holidays") == "")

ENRICH_CASES = [
    ("Physical Therapist - Home Care (PRN)",
     "Licensed Physical Therapist. Doctorate in Physical Therapy (DPT) and BLS certification. "
     "Minimum 3 years of experience. PRN schedule. Salary: $85,000 - $105,000 per year. "
     "Skills: Epic, wound care, telemetry.",
     {"employment_type": "PRN", "job_category": "Healthcare",
      "education_type": "Doctorate", "education_stream": "Allied Health",
      "currency": "USD", "min_salary": "85000"}),
    ("Senior Backend Engineer (Full-time)",
     "5+ years of experience with Python, Django, PostgreSQL, AWS and Kubernetes. "
     "Bachelor's degree in Computer Science required.",
     {"seniority_level": "Senior", "job_category": "Information Technology",
      "education_stream": "Computer Science", "employment_type": "Full Time",
      "years_of_experience_min": "5"}),
    ("Mechanical Maintenance Engineer",
     "Diploma in Mechanical Engineering. 2-4 years experience. AutoCAD and SolidWorks. Full time.",
     {"job_category": "Engineering", "education_type": "Diploma",
      "education_stream": "Engineering", "years_of_experience_max": "4"}),
    ("Registered Nurse Inpatient 2 - Pediatric Acute Care",
     "BSN preferred. Current RN licensure required. BLS certification. Night shift, full time.",
     {"job_category": "Healthcare", "education_stream": "Nursing"}),
]
bad = []
for _title, _desc, _expect in ENRICH_CASES:
    _job = {"job_title": _title, "job_description": _desc}
    enrich_job(_job)
    for _k, _want in _expect.items():
        if _job.get(_k) != _want:
            bad.append("%s: %s=%r want %r" % (_title[:20], _k, _job.get(_k), _want))
check("derived fields correct across %d job shapes" % len(ENRICH_CASES), not bad, bad[:4])

_prn = {"job_title": "Nurse (PRN)", "job_description": "PRN schedule. BSN required."}
enrich_job(_prn)
check("'PRN' not misread as the credential 'RN'",
      "RN schedule" not in (_prn.get("education_qualification") or ""),
      _prn.get("education_qualification"))

_skills = {"job_title": "Data Engineer",
           "job_description": "Python, SQL, Airflow, Snowflake, dbt and AWS required."}
enrich_job(_skills)
check("skills extracted", bool(_skills.get("skills")), _skills.get("skills"))


# ---- Workday must page past an under-reported total and enrich from detail
class WorkdayLyingTotalSession:
    """Reports total=40 while actually holding 137 postings, as VUMC does."""

    REAL_TOTAL = 137
    CLAIMED_TOTAL = 40

    def __init__(self):
        self.detail_calls = 0

    def post_json(self, url, payload, timeout=None):
        if "/wday/cxs/vumc/vumccareers/jobs" not in url:
            return None
        off, lim = payload["offset"], payload["limit"]
        items = [{"title": "Job %03d" % i,
                  "externalPath": "/job/x/Job-%03d_R-%d" % (i, i),
                  "postedOn": "Posted Today", "locationsText": "Nashville, TN"}
                 for i in range(off, min(off + lim, self.REAL_TOTAL))]
        return {"total": self.CLAIMED_TOTAL, "jobPostings": items}

    def fetch_json(self, url, timeout=None, **kw):
        self.detail_calls += 1
        return {"jobPostingInfo": {
            "jobDescription": "<p>Full time. Bachelor's degree in Nursing required. "
                              "3 years of experience. BLS. $70,000 - $90,000.</p>",
            "startDate": "2026-08-15", "endDate": "2026-09-30",
            "timeType": "Full time", "jobReqId": "R-1234"}}

    def fetch(self, *a, **kw):
        return None

    def fetch_text(self, *a, **kw):
        return None

    def head_ok(self, *a, **kw):
        return True


from job_scraper import parsers_ats as _pa
_wd_session = WorkdayLyingTotalSession()
_wd_info = _pa.workday_info_from_url("https://vumc.wd1.myworkdayjobs.com/vumccareers")
_wd_jobs = _pa.parse_workday(_wd_session, _wd_info)
check("Workday pages past an under-reported total",
      len(_wd_jobs) == WorkdayLyingTotalSession.REAL_TOTAL,
      "%d of %d" % (len(_wd_jobs), WorkdayLyingTotalSession.REAL_TOTAL))
check("Workday job details fetched", _wd_session.detail_calls > 0, _wd_session.detail_calls)
check("posted_date from detail beats 'Posted Today'",
      _wd_jobs[0]["posted_date"] == "2026-08-15", _wd_jobs[0]["posted_date"])
check("application_deadline captured",
      _wd_jobs[0]["application_deadline"] == "2026-09-30",
      _wd_jobs[0]["application_deadline"])
enrich_job(_wd_jobs[0])
check("Workday row fully enriched after detail pass",
      _wd_jobs[0]["job_category"] == "Healthcare"
      and _wd_jobs[0]["education_stream"] == "Nursing"
      and bool(_wd_jobs[0]["min_salary"]),
      {k: _wd_jobs[0][k] for k in ("job_category", "education_stream", "min_salary")})

# ------------- 11. board-URL validation, query budget, title guard
from job_scraper.ats import is_probable_board_url as _is_board
from job_scraper.company import _looks_like_job_title as _title_ok
from job_scraper import websearch as _ws

BOARD_URLS = [
    # The exact false positive seen in a live run: a "powered by iCIMS" footer
    # link to the vendor's privacy notice was dispatched as the company board.
    ("https://www.icims.com/legal/privacy-notice-website/", False),
    ("https://www.icims.com/", False),
    ("https://www.greenhouse.io/pricing", False),
    ("https://www.lever.co/blog/hiring-tips", False),
    ("https://www.personio.de/about-us/", False),
    ("https://carolinamri.icims.com/jobs/search", True),
    ("https://boards.greenhouse.io/acme", True),
    ("https://jobs.lever.co/acme", True),
    ("https://join.com/companies/acme", True),
    ("https://recruiting.paylocity.com/recruiting/jobs/All/x/Acme", True),
    ("https://vumc.wd1.myworkdayjobs.com/vumccareers", True),
]
_board_bad = [(u, _is_board(u), w) for u, w in BOARD_URLS if _is_board(u) != w]
check("vendor legal/marketing pages rejected as boards", not _board_bad, _board_bad[:3])

TITLES = [
    ("Recently Posted Jobs", False), ("Open Positions", False), ("Careers", False),
    ("Search Results", False), ("Offene Stellen", False), ("Why work with us?", False),
    ("Physical Therapist - Home Care (PRN)", True), ("Senior Backend Engineer", True),
    ("Legal Assistant (Office of Legal Affairs)", True), ("Pflegefachkraft (m/w/d)", True),
    ("Sachbearbeiter Finanzen (m/w/d)", True),
]
_title_bad = [(t, _title_ok(t), w) for t, w in TITLES if _title_ok(t) != w]
check("page headings not accepted as job titles", not _title_bad, _title_bad[:3])


# ---- search-query budget: a working website must cost zero queries
_counted = []


def _counting_engine(session, query):
    _counted.append(query)
    return ["https://unrelated.example/page"]


BUDGET_HOME = "https://www.acme-eng.com"


def _budget_careers():
    payload = json.dumps([{"@type": "JobPosting", "title": "Engineer %d" % i,
                           "url": BUDGET_HOME + "/careers/j%d" % i,
                           "description": "Apply."} for i in range(6)])
    return ('<html><head><title>Careers at ACME Engineering</title>'
            '<script type="application/ld+json">%s</script></head>'
            '<body><main><h1>Open positions</h1></main></body></html>' % payload)


class BudgetSession:
    def __init__(self, alive=True):
        self.alive = alive

    def _body(self, url):
        if not self.alive:
            return None
        u = url.split("#")[0].rstrip("/")
        if u == BUDGET_HOME:
            return ('<html><head><title>ACME Engineering Ltd</title></head><body>'
                    '<h1>ACME Engineering</h1><a href="/careers">Careers</a></body></html>')
        if u.startswith(BUDGET_HOME + "/careers"):
            return _budget_careers()
        return None

    def fetch(self, url, timeout=None, method="GET", retries=None, **kw):
        body = self._body(url)
        return FakeResponse(url, body) if body else None

    def fetch_text(self, url, timeout=None, **kw):
        r = self.fetch(url, timeout=timeout)
        return r.text if r else None

    def fetch_json(self, url, timeout=None, **kw):
        return None

    def post_json(self, url, payload, timeout=None):
        return None

    def head_ok(self, url, timeout=12):
        return True


def _measure_queries(session):
    _counted.clear()
    _ws._QUERY_CACHE.clear()
    _ws._ENGINE_STATE.clear()
    saved = _ws._ENGINES
    _ws._ENGINES = (("counting", _counting_engine),)
    try:
        details = company.process_company_details(
            ("ACME Engineering Ltd", "www.acme-eng.com", "USA"),
            session=session, enable_search=True)
    finally:
        _ws._ENGINES = saved
    return len(_counted), details


_working_queries, _working = _measure_queries(BudgetSession(alive=True))
check("a working website costs zero search queries", _working_queries == 0,
      _working_queries)
check("...and still extracts its jobs", _working["num_positions"] == 6,
      _working["num_positions"])

_dead_queries, _dead = _measure_queries(BudgetSession(alive=False))
# A company whose site is dead (or absent) is found by NAME, which is search
# only, so it gets the larger no-website budget rather than the default 5.
check("a dead website escalates to the no-website budget (%d)"
      % config.SEARCH_BUDGET_NO_WEBSITE,
      _dead_queries <= config.SEARCH_BUDGET_NO_WEBSITE, _dead_queries)
check("...and it spends more than the working-site budget rather than giving up",
      _dead_queries > config.MAX_SEARCH_QUERIES_PER_COMPANY, _dead_queries)

# ---------------- 12. cost containment: asset URLs and render retries
# A live run spent ~18 minutes of browser time on one iCIMS download servlet,
# rendered 6 times at a 90s timeout, twice, blocking both render workers.
ASSET_URLS = [
    ("https://mqimaging.icims.com/icims2/servlet/icims2"
     "?module=AppInert&action=download&id=57&hashed=227258933", False),
    ("https://acme.icims.com/icims2/servlet/icims2?module=AppJob", False),
    ("https://acme.workable.com/assets/logo.png", False),
    ("https://acme.teamtailor.com/jobs/brochure.pdf", False),
    ("https://acme.csod.com/media/download?id=9", False),
    ("https://acme.icims.com/careers/job/12345", True),
    ("https://acme.recruitee.com/o/engineer", True),
    ("https://boards.greenhouse.io/acme/jobs/456", True),
]
_asset_bad = [(u, _is_board(u), w) for u, w in ASSET_URLS if _is_board(u) != w]
check("download/servlet/asset URLs rejected as boards", not _asset_bad, _asset_bad[:3])

import job_scraper.browser_fetch as _bf

check("render attempts bounded", config.BROWSER_RENDER_MAX_ATTEMPTS <= 2,
      config.BROWSER_RENDER_MAX_ATTEMPTS)
check("probe timeout bounded", config.PROBE_TIMEOUT <= 20, config.PROBE_TIMEOUT)


class _TimeoutWorker:
    calls = 0
    alive = True

    def submit(self, request):
        _TimeoutWorker.calls += 1
        request.error = Exception(
            "Page.goto: Timeout 45000ms exceeded.\nCall log:\n  - navigating")
        request.done.set()
        return True


_saved_ensure = _bf._ensure_workers
_bf._ensure_workers = lambda: [_TimeoutWorker()]
_bf._BROWSER_UNAVAILABLE.clear()
_TimeoutWorker.calls = 0
_ASSET_URL = ("https://mqimaging.icims.com/icims2/servlet/icims2"
              "?module=AppInert&action=download")
_first = _bf.fetch_rendered_html(_ASSET_URL)
_attempts_after_first = _TimeoutWorker.calls
_second = _bf.fetch_rendered_html(_ASSET_URL)
_bf._ensure_workers = _saved_ensure

check("a navigation timeout is not retried", _attempts_after_first == 1,
      _attempts_after_first)
check("an unrenderable URL is not rendered again",
      _TimeoutWorker.calls == _attempts_after_first, _TimeoutWorker.calls)
check("failed render returns None", _first is None and _second is None)

# ------------------- 13. expanded vendor coverage + JSON board prober
import re as _re
import job_scraper.ats as _ats_mod
from job_scraper import parsers_ats as _pa
from job_scraper.company import _KNOWN_GENERIC as _GENERIC, _API_PARSERS as _APIS

_vendor_names = set(_re.findall(r'^    \("([a-z0-9_]+)", re',
                                open(_ats_mod.__file__, encoding="utf-8").read(), _re.M))
check("vendor coverage is broad", len(_vendor_names) >= 120, len(_vendor_names))
_unrouted = sorted(_vendor_names - (set(_GENERIC) | set(_APIS)))
check("every detected vendor has an extraction route", not _unrouted, _unrouted)

REGIONAL_URLS = [
    ("https://acme.gupy.io/", "gupy"),
    ("https://acme.kenoby.com/vagas", "kenoby"),
    ("https://acme.solides.jobs/vagas", "solides"),
    ("https://99jobs.com/acme", "99jobs"),
    ("https://acme.dvinci-hr.com/de/jobs", "dvinci"),
    ("https://acme.rexx-systems.com/stellen", "rexx"),
    ("https://www.interamt.de/koop/app/stelle?id=123", "interamt"),
    ("https://acme.greythr.com/careers", "greythr"),
    ("https://acme.turbohire.co/jobs", "turbohire"),
    ("https://acme.livehire.com/careers", "livehire"),
    ("https://acme.pageuppeople.com/jobs", "pageup"),
    ("https://acme.elmotalent.com.au/careers", "elmo"),
    ("https://acme.taleez.com/jobs", "taleez"),
    ("https://acme.catsone.com/careers/jobs", "catsone"),
    ("https://acme.applicantpro.com/jobs/", "applicantpro"),
    ("https://acme.recruiterbox.com/jobs", "trakstar"),
    # GovernmentJobs hosts thousands of employers, so it is an aggregator, not
    # one company's board -- correctly excluded even though it is detected.
    ("https://www.governmentjobs.com/careers/richmondva", None),
    ("https://acme.peopleadmin.com/postings/search", "peopleadmin"),
    ("https://www.gupy.io/", None),
    ("https://www.livehire.com/pricing", None),
]


def _resolve_vendor(url):
    name, _ = _detect(url)
    return name if (name and _is_board(url)) else None


_regional_bad = [(u, _resolve_vendor(u), w)
                 for u, w in REGIONAL_URLS if _resolve_vendor(u) != w]
check("regional vendor URLs resolve (BR/DACH/IN/AU/public sector)",
      not _regional_bad, _regional_bad[:3])


class _JsonBoardSession:
    """Serves one JSON payload at <board>/jobs.json and nothing else."""

    def __init__(self, payload):
        self.payload = payload
        self.tried = []

    def fetch_json(self, url, timeout=None, **kw):
        self.tried.append(url)
        return self.payload if url.endswith("/jobs.json") else None


JSON_SHAPES = {
    "flat list": ([{"title": "Engenheiro de Software", "url": "https://x/1",
                    "city": "Sao Paulo", "descricao": "<p>Aplique</p>",
                    "dataPublicacao": "2026-08-01"},
                   {"title": "Analista de Dados", "url": "https://x/2"},
                   {"title": "Gerente de Produto", "url": "https://x/3"}], 3),
    "under data": ({"data": [{"name": "Pflegefachkraft (m/w/d)", "href": "/job/1",
                              "ort": "Berlin"},
                             {"name": "Sachbearbeiter", "href": "/job/2"},
                             {"name": "Erzieher", "href": "/job/3"}]}, 3),
    "nested results": ({"payload": {"results": [
        {"jobTitle": "Nurse", "applyUrl": "https://x/a", "locationName": "Sydney"},
        {"jobTitle": "Doctor", "applyUrl": "https://x/b"},
        {"jobTitle": "Cleaner", "applyUrl": "https://x/c"}]}}, 3),
    "vagas key": ({"vagas": [{"titulo": "Vendedor", "permalink": "https://x/v1"},
                             {"titulo": "Caixa", "permalink": "https://x/v2"},
                             {"titulo": "Estoquista", "permalink": "https://x/v3"}]}, 3),
    "single item": ({"jobs": [{"title": "Only One", "url": "https://x/1"}]}, 0),
    "not postings": ({"data": [{"id": 1, "colour": "red"}, {"id": 2, "colour": "blue"},
                               {"id": 3, "colour": "green"}]}, 0),
    "empty": ({"jobs": []}, 0),
}
_json_bad = []
for _label, (_payload, _want) in JSON_SHAPES.items():
    _sess = _JsonBoardSession(_payload)
    _found = _pa.probe_json_board(_sess, "https://acme.unknownats.com/careers", "unknownats")
    if len(_found) != _want:
        _json_bad.append("%s: %d != %d" % (_label, len(_found), _want))
check("JSON prober reads unfamiliar payload shapes and rejects noise",
      not _json_bad, _json_bad)

_probe_sess = _JsonBoardSession(JSON_SHAPES["flat list"][0])
_probe_jobs = _pa.probe_json_board(_probe_sess, "https://acme.unknownats.com/careers", "gupy")
check("JSON prober populates location/date/description",
      bool(_probe_jobs and _probe_jobs[0]["job_location"]
           and _probe_jobs[0]["posted_date"] and _probe_jobs[0]["job_description"]))
check("JSON prober tries at most 4 endpoints", len(_probe_sess.tried) <= 4,
      len(_probe_sess.tried))

# --------- 14. alias normalisation, aggregator exclusion, unlisted vendors
check("vendor coverage expanded", len(_vendor_names) >= 230, len(_vendor_names))

ALIAS_CASES = [
    ("Oracle Taleo", "taleo"), ("SAP SuccessFactors", "successfactors"),
    ("SAP Success Factors", "successfactors"), ("Infinite Talent", "brassring"),
    ("Kenexa", "brassring"), ("KeldairHR", "keldair"), ("Hyrell", "keldair"),
    ("Rival", "silkroad"), ("SilkRoad Technology", "silkroad"),
    ("SENTIO", "sprockets"), ("Greenhouse Recruiting", "greenhouse"),
    ("monday.com", "jobflows"), ("Recruitment CRM", "recruitcrm"),
    ("Recruiterbox", "trakstar"), ("Mitratech Trakstar", "trakstar"),
    ("Ceridian", "dayforce"), ("SchoolSpring", "applitrack"),
    ("UltiPro", "ukg"), ("Kronos", "ukg"), ("Access Vincere Evo", "vincere"),
]
_alias_bad = [(r, _ats_mod.canonical_ats_name(r), w)
              for r, w in ALIAS_CASES if _ats_mod.canonical_ats_name(r) != w]
check("vendor renames normalise to one canonical name", not _alias_bad, _alias_bad[:3])

AGGREGATORS = [
    "https://www.indeed.com/jobs?q=x",
    "https://www.glassdoor.com/Job/x-SRCH.htm",
    "https://www.linkedin.com/jobs/view/123",
    "https://www.teamworkonline.com/jobs",
    "https://www.vagas.com.br/cargo/analista",
    "https://www.governmentjobs.com/careers/richmondva",
]
check("aggregators are never treated as one company's board",
      not [u for u in AGGREGATORS if _is_board(u)],
      [u for u in AGGREGATORS if _is_board(u)][:2])
REAL_BOARDS = [
    "https://boards.greenhouse.io/acme",
    "https://acme.applitrack.com/acme/onlineapp/",
    "https://acme.comeet.com/jobs",
    "https://acme.gohire.io/jobs",
]
check("real boards still accepted",
      all(_is_board(u) for u in REAL_BOARDS),
      [u for u in REAL_BOARDS if not _is_board(u)])

_BOARD_HTML = ('<html><head><title>Open Roles at Acme</title></head><body><ul>'
               + "".join('<li class="job-item"><a href="/jobs/%d-eng">Engineer %02d</a></li>'
                         % (i, i) for i in range(9))
               + '</ul><p>Apply now</p></body></html>')
check("structural detector recognises a board",
      _ats_mod.looks_like_job_board(_BOARD_HTML))
check("structural detector rejects an about page",
      not _ats_mod.looks_like_job_board(
          '<html><head><title>About Acme</title></head><body>'
          '<p>We make widgets since 1990.</p></body></html>'))
check("structural detector rejects a news page",
      not _ats_mod.looks_like_job_board(
          '<html><head><title>Acme News</title></head><body>'
          '<a href="/news/1">Story</a><a href="/news/2">Story</a></body></html>'))


# ---- a career page handing off to a vendor with NO pattern at all
UNLISTED_HOME = "https://www.widgetco.example"
UNLISTED_BOARD = "https://widgetco.brand-new-ats-nobody-knows.example/vacancies"
UNLISTED_N = 14


class UnlistedVendorSession:
    def _body(self, url):
        u = url.split("#")[0].rstrip("/")
        if u == UNLISTED_HOME:
            return ('<html><head><title>WidgetCo Ltd</title></head><body>'
                    '<h1>WidgetCo</h1><a href="/careers">Careers</a></body></html>')
        if u == UNLISTED_HOME + "/careers":
            return ('<html><head><title>Careers at WidgetCo</title></head><body><main>'
                    '<h1>Work with us</h1><p>See our current openings</p>'
                    '<a href="%s">View all vacancies</a></main></body></html>' % UNLISTED_BOARD)
        if u == UNLISTED_BOARD.rstrip("/"):
            return ('<html><head><title>Open Roles at WidgetCo</title></head><body><ul>'
                    + "".join('<li class="job-item"><a href="/jobs/%d-eng">Engineer %02d</a></li>'
                              % (i, i) for i in range(UNLISTED_N))
                    + '</ul><p>Apply now</p></body></html>')
        return None

    def fetch(self, url, timeout=None, method="GET", retries=None, **kw):
        body = self._body(url)
        return FakeResponse(url, body) if body else None

    def fetch_text(self, url, timeout=None, **kw):
        r = self.fetch(url, timeout=timeout)
        return r.text if r else None

    def fetch_json(self, url, timeout=None, **kw):
        if url.startswith(UNLISTED_BOARD.rstrip("/")) and url.endswith("/jobs.json"):
            return {"jobs": [{"title": "Engineer %02d" % i,
                              "url": UNLISTED_BOARD + "/j%d" % i,
                              "location": "Leeds",
                              "description": "Apply now. Responsibilities."}
                             for i in range(UNLISTED_N)]}
        return None

    def post_json(self, url, payload, timeout=None):
        return None

    def head_ok(self, url, timeout=12):
        return True


check("test premise: that vendor really is unlisted",
      _detect(UNLISTED_BOARD)[0] is None, _detect(UNLISTED_BOARD)[0])
_unlisted = company.process_company_details(
    ("WidgetCo Ltd", UNLISTED_HOME, "UK"),
    session=UnlistedVendorSession(), enable_search=False)
check("a vendor with no pattern is still extracted",
      _unlisted["num_positions"] == UNLISTED_N, _unlisted["num_positions"])
check("structurally-detected rows are attributed as such",
      _unlisted["jobs"] and _unlisted["jobs"][0]["source"] == "structural-board-detection",
      _unlisted["jobs"][0]["source"] if _unlisted["jobs"] else None)

# --------------- 15. cost bugs found in a 2,000-company live run
from job_scraper.parsers_ats import _json_candidate_urls as _json_urls

# "<origin>.json" is not a hostname: the probe generated
# https://deliverymuch.gupy.io.json and paid a DNS failure plus retries.
_bad_json = [u for board in ("https://deliverymuch.gupy.io",
                             "https://cdn.phenompeople.com")
             for u in _json_urls(board)
             if u.rstrip("/").endswith((".io.json", ".com.json", ".br.json"))]
check("JSON probe never builds an invalid '<host>.json'", not _bad_json, _bad_json)
check("JSON probe still uses '<path>.json' when a path exists",
      any(u.endswith("/careers.json")
          for u in _json_urls("https://acme.unknownats.com/careers")))

CDN_URLS = [
    ("https://cdn.phenompeople.com", False),
    ("https://static.workable.com/widget.js", False),
    ("https://assets.greenhouse.io/logo.png", False),
    ("https://boards.greenhouse.io/acme", True),
    ("https://deliverymuch.gupy.io", True),
]
_cdn_bad = [(u, _is_board(u), w) for u, w in CDN_URLS if _is_board(u) != w]
check("CDN/asset hosts rejected as boards", not _cdn_bad, _cdn_bad[:3])


# ---- one external board, reached from every common career path
DM_HOME = "https://www.deliverymuch.com.br"
DM_BOARD = "https://deliverymuch.gupy.io"
DM_N = 9
DM_PATHS = ["/careers", "/career", "/jobs", "/jobs/careers", "/careers/jobs",
            "/join-us", "/career-opportunities", "/work-with-us", "/vacancies",
            "/join-our-team", "/about/careers", "/career/jobs", "/karriere",
            "/stellenangebote", "/job-openings", "/careers-at", "/recruiting"]


class DuplicateBoardSession:
    """Every career path links to the same external board (as observed live)."""

    def __init__(self):
        self.json_probes = 0

    def _body(self, url):
        u = url.split("#")[0].rstrip("/")
        if u == DM_HOME:
            return ('<html><head><title>Delivery Much</title></head><body>'
                    '<a href="/careers">Carreiras</a></body></html>')
        if any(u == DM_HOME + p for p in DM_PATHS):
            return ('<html><head><title>Carreiras</title></head><body><main>'
                    '<h1>Vagas</h1><a href="%s">Ver vagas</a></main></body></html>'
                    % DM_BOARD)
        if u.startswith(DM_BOARD):
            return '<html><body>Gupy board shell</body></html>'
        return None

    def fetch(self, url, timeout=None, method="GET", retries=None, **kw):
        body = self._body(url)
        return FakeResponse(url, body) if body else None

    def fetch_text(self, url, timeout=None, **kw):
        r = self.fetch(url, timeout=timeout)
        return r.text if r else None

    def fetch_json(self, url, timeout=None, **kw):
        if url.startswith(DM_BOARD):
            self.json_probes += 1
            if url.endswith("/jobs.json"):
                return {"jobs": [{"title": "Vaga %02d" % i,
                                  "url": DM_BOARD + "/job/%d" % i,
                                  "city": "Blumenau", "descricao": "Aplique"}
                                 for i in range(DM_N)]}
        return None

    def post_json(self, url, payload, timeout=None):
        return None

    def head_ok(self, url, timeout=12):
        return True


_dm_session = DuplicateBoardSession()
_dm = company.process_company_details(
    ("DELIVERY MUCH TECNOLOGIA SA", DM_HOME, "Brazil"),
    session=_dm_session, enable_search=False)
check("positions deduped across %d career paths" % len(DM_PATHS),
      _dm["num_positions"] == DM_N, _dm["num_positions"])
check("one external board is processed once, not once per path",
      _dm_session.json_probes <= 4, _dm_session.json_probes)

# ---------------- 16. deployment hardening (long unattended VM run)
import glob as _glob
import logging as _logging
import time as _time
from logging.handlers import RotatingFileHandler as _RFH
import job_scraper.session as _se
import job_scraper.browser_fetch as _bfmod

check("per-company time budget configured",
      config.COMPANY_TIME_BUDGET > 0, config.COMPANY_TIME_BUDGET)
check("ATS detail enrichment bounded (sequential fetches)",
      config.MAX_ATS_DETAIL_ENRICH <= 1000, config.MAX_ATS_DETAIL_ENRICH)
check("log rotation configured", config.LOG_MAX_BYTES > 0 and config.LOG_BACKUP_COUNT > 0,
      (config.LOG_MAX_BYTES, config.LOG_BACKUP_COUNT))
check("a rotating handler is installed, not a plain FileHandler",
      any(isinstance(h, _RFH) for h in _logging.getLogger().handlers)
      or not _logging.getLogger().handlers,
      [type(h).__name__ for h in _logging.getLogger().handlers])

# run() twice must not stack handlers
_handlers_before = len(_logging.getLogger().handlers)
_dupe_in = os.path.join(tmpdir, "dupe.xlsx")
_wb = openpyxl.Workbook(); _ws = _wb.active
_ws.append(["KEYID", "COMPANY_NAME", "COUNTRY", "ENTITY_TYPE", "WEBSITE"])
_ws.append(["k", "Nowhere GmbH", "Germany", "Independent", ""])
_wb.save(_dupe_in)


class _DeadSession:
    def fetch(self, *a, **kw): return None
    def fetch_text(self, *a, **kw): return None
    def fetch_json(self, *a, **kw): return None
    def post_json(self, *a, **kw): return None
    def head_ok(self, *a, **kw): return False


pipeline.ScrapeSession = _DeadSession
pipeline.run(input_file=_dupe_in, output_file=os.path.join(tmpdir, "d1.csv"),
             workers=1, resume=False, enable_search=False, quiet=True)
pipeline.run(input_file=_dupe_in, output_file=os.path.join(tmpdir, "d2.csv"),
             workers=1, resume=False, enable_search=False, quiet=True)
check("run() twice does not duplicate log handlers",
      len(_logging.getLogger().handlers) == _handlers_before,
      (_handlers_before, len(_logging.getLogger().handlers)))

# remove_retryable_rows must stream a large file, not load it
_big = os.path.join(tmpdir, "big_output.csv")
with io.open(_big, "w", encoding="utf-8-sig", newline="") as _f:
    _w = csv.DictWriter(_f, fieldnames=config.OUTPUT_COLUMNS, extrasaction="ignore")
    _w.writeheader()
    for _i in range(4000):
        _w.writerow({"company_name": "Co%d" % _i,
                     "job_status": "Unreachable" if _i % 2 else "Active",
                     "job_description": "x" * 400, "num_positions": 0})
_removed = pipeline.remove_retryable_rows(_big)
with io.open(_big, encoding="utf-8-sig") as _f:
    _kept = sum(1 for _ in _f) - 1
check("retry purge removes exactly the transient rows",
      _removed == 2000 and _kept == 2000, (_removed, _kept))

# host-failure memos bounded across many hosts
_se.reset_host_failures()
for _i in range(_se._HOST_MEMO_LIMIT + 5000):
    _se._mark_host_failure("h%d.example" % _i)
check("host-failure memo stays bounded over many hosts",
      len(_se._DEAD_HOSTS) <= _se._HOST_MEMO_LIMIT + 1, len(_se._DEAD_HOSTS))

# rendered-page cache prunes expired entries and respects a size cap
_cache_dir = os.path.join(tmpdir, "browser_cache")
os.makedirs(_cache_dir, exist_ok=True)
_saved_dir, _saved_cap = _bfmod._CACHE_DIR, _bfmod._CACHE_MAX_BYTES
_bfmod._CACHE_DIR, _bfmod._CACHE_MAX_BYTES = _cache_dir, 50000
for _i in range(30):
    _path = os.path.join(_cache_dir, "f%02d.html" % _i)
    io.open(_path, "w").write("x" * 5000)
    os.utime(_path, (_time.time() - _i * 10, _time.time() - _i * 10))
for _i in range(5):
    _path = os.path.join(_cache_dir, "expired%d.html" % _i)
    io.open(_path, "w").write("x" * 5000)
    _old = _time.time() - _bfmod._CACHE_TTL_SECONDS - 100
    os.utime(_path, (_old, _old))
_bfmod._prune_cache()
_remaining = os.listdir(_cache_dir)
_bytes = sum(os.path.getsize(os.path.join(_cache_dir, f)) for f in _remaining)
check("expired rendered pages are deleted",
      not [f for f in _remaining if f.startswith("expired")])
check("rendered-page cache respects its size cap", _bytes <= 50000, _bytes)
_bfmod._CACHE_DIR, _bfmod._CACHE_MAX_BYTES = _saved_dir, _saved_cap

check("browser context recycling configured",
      _bfmod._CONTEXT_RECYCLE_AFTER > 0, _bfmod._CONTEXT_RECYCLE_AFTER)


# ------------- 18. the input WEBSITE column is not a dependency
# 234,685 of the 405,210 input rows have no WEBSITE at all.  A company must be
# findable from its NAME alone: guessed domains first (free), then search, then
# an ATS board published under its name.

_GUESS = company._guess_hosts

_de_guesses = _GUESS("Acme Foods Ltd", "Germany")
check("guessed domains cover the country TLD and .com",
      {"acmefoods.de", "acmefoods.com", "acme-foods.de"}.issubset(set(_de_guesses)),
      _de_guesses)
_br_guesses = _GUESS("Vortice Industrial Ltda", "Brazil")
check("Brazilian companies are guessed at .com.br",
      "vorticeindustrial.com.br" in _br_guesses, _br_guesses)
check("a short one-word name is too generic to guess from",
      _GUESS("Star", "USA") == [], _GUESS("Star", "USA"))
check("guessing respects its limit (%d)" % config.DOMAIN_GUESS_LIMIT,
      len(_GUESS("Alpha Beta Gamma Delta Trading", "India")) <= config.DOMAIN_GUESS_LIMIT,
      len(_GUESS("Alpha Beta Gamma Delta Trading", "India")))

_order_br = websearch._group_order("Brazil")
_order_de = websearch._group_order("Germany")
check("Brazilian companies are asked about Gupy/Solides first",
      _order_br[0] == 3, _order_br)
check("German and Indian companies are asked about their vendors first",
      _order_de[0] == 4, _order_de)
check("US companies keep the default vendor order",
      websearch._group_order("USA")[0] == 0)
check("no vendor group is unreachable",
      sorted(_order_br) == list(range(len(websearch._ATS_SITE_GROUPS))), _order_br)

websearch.set_query_budget(config.MAX_SEARCH_QUERIES_PER_COMPANY)
websearch._consume_budget()
websearch._consume_budget()
websearch.raise_query_budget(config.SEARCH_BUDGET_NO_WEBSITE)
check("the no-website budget lifts the remaining queries",
      websearch.query_budget_left() == config.SEARCH_BUDGET_NO_WEBSITE,
      websearch.query_budget_left())
websearch.raise_query_budget(1)
check("...and never lowers a budget",
      websearch.query_budget_left() == config.SEARCH_BUDGET_NO_WEBSITE,
      websearch.query_budget_left())


def _mini_careers(host, company_name, count):
    payload = json.dumps([
        {"@context": "https://schema.org", "@type": "JobPosting",
         "title": "Operator %d" % i,
         "url": "%s/karriere/job/%d" % (host, i),
         "datePosted": "2026-08-01",
         "employmentType": "FULL_TIME",
         "hiringOrganization": {"@type": "Organization", "name": company_name},
         "description": "<p>Apply now. Responsibilities and qualifications.</p>"}
        for i in range(count)])
    return ('<html><head><title>Karriere bei %s</title>'
            '<script type="application/ld+json">%s</script></head><body>'
            '<h1>Offene Stellen</h1><p>Apply now. Job description, '
            'responsibilities, qualifications.</p></body></html>'
            % (company_name, payload))


class NoWebsiteSession:
    """Serves ONE host and nothing else -- as if only a guessed domain existed."""

    def __init__(self, host, company_name, jobs=3, homepage_mentions=True):
        self.host = host
        self.company_name = company_name
        self.jobs = jobs
        self.mentions = homepage_mentions
        self.fetched = []

    def _body(self, url):
        u = url.split("#")[0].rstrip("/")
        if not u.startswith(self.host):
            return None
        if u == self.host:
            title = self.company_name if self.mentions else "Domain For Sale Inc"
            body = ("<h1>%s</h1><a href=\"/karriere\">Karriere</a>" % title
                    if self.mentions else
                    "<h1>Cheap Widgets Online</h1><p>Buy widgets. Nothing else here. "
                    + ("filler " * 400) + "</p>")
            return "<html><head><title>%s</title></head><body>%s</body></html>" % (title, body)
        if u.endswith("/karriere"):
            return _mini_careers(self.host, self.company_name, self.jobs)
        return None

    def fetch(self, url, timeout=None, method="GET", retries=None, **kw):
        self.fetched.append(url)
        body = self._body(url)
        return FakeResponse(url, body) if body else None

    def fetch_text(self, url, timeout=None, **kw):
        r = self.fetch(url, timeout=timeout)
        return r.text if r else None

    def fetch_json(self, url, timeout=None, **kw):
        return None

    def post_json(self, url, payload, timeout=None):
        return None

    def head_ok(self, url, timeout=12):
        return self._body(url) is not None


_saved_resolves = company._resolves

# --- a company with an EMPTY website column, found purely by guessing -------
company._resolves = lambda host: host == "guessablefoods.de"
_guess_session = NoWebsiteSession("https://guessablefoods.de", "Guessable Foods GmbH", jobs=3)
_guess_details = company.process_company_details(
    ("Guessable Foods GmbH", "", "Germany"), session=_guess_session,
    enable_search=False)          # enable_search=False: zero search queries
check("a company with no website is found by guessing its domain",
      _guess_details["status"] == "ok", _guess_details["status"])
check("...and all of its postings are returned",
      _guess_details["num_positions"] == 3, _guess_details["num_positions"])
check("...with the guess recorded in website_discovery",
      _guess_details["website_discovery"] == "name_domain_guess",
      _guess_details["website_discovery"])
check("...and the resolved address recorded for audit",
      _guess_details["resolved_website"] == "https://guessablefoods.de",
      _guess_details["resolved_website"])

# --- a guessed domain that belongs to someone else must be rejected ---------
company._resolves = lambda host: host == "zephyr.com"
_wrong_session = NoWebsiteSession("https://zephyr.com", "Zephyr Analytics Ltd",
                                  jobs=3, homepage_mentions=False)
_wrong_details = company.process_company_details(
    ("Zephyr Analytics Ltd", "", "USA"), session=_wrong_session, enable_search=False)
check("a guessed domain with no evidence of the company is not accepted",
      _wrong_details["website_discovery"] == "none",
      _wrong_details["website_discovery"])
check("...so no other company's jobs are attributed to it",
      _wrong_details["num_positions"] == 0, _wrong_details["num_positions"])

# --- no website, no guessable domain: found via an ATS board by name --------
GH_BOARD = "https://boards.greenhouse.io/boardonlysystems"
GH_API = "https://boards-api.greenhouse.io/v1/boards/boardonlysystems/jobs?content=true"


def _board_engine(session, query):
    # Only the ATS site: query knows anything; the "official website" query
    # returns nothing, exactly like a company with no site of its own.
    if "site:" in query and "greenhouse" in query:
        return [GH_BOARD]
    return []


class BoardOnlySession:
    def __init__(self):
        self.json_calls = []

    def fetch(self, url, timeout=None, method="GET", retries=None, **kw):
        if url.split("#")[0].rstrip("/") == GH_BOARD:
            return FakeResponse(url, "<html><body>Board</body></html>")
        return None

    def fetch_text(self, url, timeout=None, **kw):
        r = self.fetch(url, timeout=timeout)
        return r.text if r else None

    def fetch_json(self, url, timeout=None, **kw):
        self.json_calls.append(url)
        if url.startswith("https://boards-api.greenhouse.io/v1/boards/boardonlysystems/jobs"):
            return {"jobs": [
                {"title": "Field Technician", "absolute_url": GH_BOARD + "/jobs/1",
                 "updated_at": "2026-08-01T00:00:00Z", "location": {"name": "Austin, TX"},
                 "content": "<p>Apply now. Responsibilities.</p>"},
                {"title": "Dispatcher", "absolute_url": GH_BOARD + "/jobs/2",
                 "updated_at": "2026-08-02T00:00:00Z", "location": {"name": "Austin, TX"},
                 "content": "<p>Apply now. Qualifications.</p>"},
            ]}
        return None

    def post_json(self, url, payload, timeout=None):
        return None

    def head_ok(self, url, timeout=12):
        return True


company._resolves = lambda host: False        # nothing guessable resolves
_saved_engines = websearch._ENGINES
websearch._QUERY_CACHE.clear()
websearch._ENGINE_STATE.clear()
websearch._ENGINES = (("board", _board_engine),)
try:
    _board_details = company.process_company_details(
        ("Boardonly Systems", "", "USA"), session=BoardOnlySession(),
        enable_search=True)
finally:
    websearch._ENGINES = _saved_engines
    company._resolves = _saved_resolves

check("a company with no website is still found through its ATS board",
      _board_details["status"] == "ok", _board_details["status"])
check("...and its postings are counted",
      _board_details["num_positions"] == 2, _board_details["num_positions"])
check("...recorded as an ATS-board-by-name discovery",
      _board_details["career_page_discovery_method"].startswith("ats_board_name_search"),
      _board_details["career_page_discovery_method"])

shutil.rmtree(tmpdir, ignore_errors=True)
print("\n%d passed, %d failed" % (len(PASS), len(FAIL)))
if FAIL:
    print("FAILED: " + ", ".join(FAIL))
sys.exit(1 if FAIL else 0)
